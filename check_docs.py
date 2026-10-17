"""Keep the sweep docs consistent across regions (added 2026-10-10, Alexandra: "no deviating").

One rulebook (CLAUDE.md) holds every rule and the one sweep procedure. Each region has a source list in
sources/<region>.md, all in the same fixed format. This script fails (exit 1) when:
  - a sources file is missing a header field, uses a section that is not allowed, or has no numbered sources;
  - a file's "Source count" does not match the numbered entries under "## Sources", or numbers repeat;
  - a region has a sources file but no row in CLAUDE.md's Regions table (or the reverse);
  - CLAUDE.md is missing the shared sweep / processing / regions sections.
A sources file may hold SOURCES ONLY: rules belong in CLAUDE.md, where they apply to every region.

Usage:
  python check_docs.py              # check the real files
  python check_docs.py --self-test  # prove the checks can fail
  python check_docs.py --attestation <review.xlsx> [--region marin|napa|telluride]
      # completion check on a sweep workbook (added 2026-10-10): its "Attestation Log" must have one row per
      # source in sources/<region>.md, every row with a Result, and no banned result such as "not reached".
      # Region is taken from the file name (daily_sweep = marin, napa_sweep, telluride_sweep) unless given.
"""
import os, re, sys, tempfile, shutil

HERE = os.path.dirname(os.path.abspath(__file__))
FIELDS = ['Region key', '`county` value on records', 'Sweep window', 'Review workbook', 'Source count']
ALLOWED_H2 = {'Sources', 'Not yet verified', 'Retired or seasonal', 'Reference'}
BANNED_RESULTS = ['recurring already in db', 'not reached', 'unverified', "didn't get to", 'did not get to', 'not checked']
CLAUDE_SECTIONS = ['## Regions (Marin, Napa, Telluride+)', '## Weekly sweep (one procedure for every region)',
                   '## Processing an approved sweep (every region)']


def check(root):
    errs = []
    src_dir = os.path.join(root, 'sources')
    claude = open(os.path.join(root, 'CLAUDE.md'), encoding='utf-8').read()
    for sec in CLAUDE_SECTIONS:
        if sec not in claude:
            errs.append(f'CLAUDE.md: missing section "{sec}"')
    reg = claude[claude.find('## Regions'):]
    reg = reg[:reg.find('\n## ', 5)] if '\n## ' in reg[5:] else reg
    table_files = set(re.findall(r'`sources/([a-z]+)\.md`', reg))
    files = sorted(f for f in os.listdir(src_dir) if f.endswith('.md')) if os.path.isdir(src_dir) else []
    if not files:
        errs.append('sources/: no region files')
    for f in files:
        key = f[:-3]
        t = open(os.path.join(src_dir, f), encoding='utf-8').read()
        tag = f'sources/{f}'
        if not t.startswith('# Sweep sources: '):
            errs.append(f'{tag}: first line must be "# Sweep sources: <Region>"')
        vals = {}
        for fld in FIELDS:
            m = re.search(r'^\| ' + re.escape(fld) + r' \| (.+?) \|\s*$', t, re.M)
            if not m:
                errs.append(f'{tag}: header field "{fld}" missing')
            else:
                vals[fld] = m.group(1).strip()
        if vals.get('Region key') and vals['Region key'].strip('`') != key:
            errs.append(f'{tag}: Region key {vals["Region key"]} does not match file name')
        for h in re.findall(r'^## (.+)$', t, re.M):
            if h.strip() not in ALLOWED_H2:
                errs.append(f'{tag}: section "## {h}" not allowed (allowed: {", ".join(sorted(ALLOWED_H2))}). Rules go in CLAUDE.md.')
        m = re.search(r'^## Sources\s*$(.*?)(?=^## |\Z)', t, re.M | re.S)
        if not m:
            errs.append(f'{tag}: no "## Sources" section')
            continue
        nums = [int(n) for n in re.findall(r'^(\d+)\. \*\*', m.group(1), re.M)]
        if not nums:
            errs.append(f'{tag}: no numbered sources ("N. **Name** — ...")')
        if len(nums) != len(set(nums)):
            errs.append(f'{tag}: repeated source numbers {sorted(n for n in set(nums) if nums.count(n) > 1)}')
        cnt = vals.get('Source count', '')
        if not cnt.isdigit():
            errs.append(f'{tag}: Source count must be a number, got "{cnt}"')
        elif int(cnt) != len(nums):
            errs.append(f'{tag}: Source count says {cnt} but {len(nums)} numbered sources found')
        if key not in table_files:
            errs.append(f'{tag}: no row for it in CLAUDE.md "Regions" table')
    for k in table_files - {f[:-3] for f in files}:
        errs.append(f'CLAUDE.md Regions table names sources/{k}.md, which does not exist '
                    f'(if you downloaded these files, that download probably failed: download sources/{k}.md again)')
    return errs


def region_from_name(path):
    n = os.path.basename(path).lower()
    for pre, key in [('daily_sweep', 'marin'), ('napa_sweep', 'napa'), ('telluride', 'telluride')]:
        if n.startswith(pre):
            return key
    return None


def check_attestation(xlsx, region=None, root=HERE, partial=False):
    """Completion check of one review workbook: the Sources Used sheet, and (unless partial, i.e. a one-source review)
    one Attestation Log row per source in the region file."""
    import openpyxl
    region = region or region_from_name(xlsx)
    if not region:
        return [f'cannot tell the region from "{os.path.basename(xlsx)}"; pass --region']
    src = os.path.join(root, 'sources', f'{region}.md')
    m = re.search(r'^\| Source count \| (\d+) \|', open(src, encoding='utf-8').read(), re.M)
    expected = int(m.group(1))
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if 'Attestation Log' not in wb.sheetnames:
        return ['workbook has no "Attestation Log" sheet']
    rows = [r for r in wb['Attestation Log'].iter_rows(values_only=True)]
    head = [str(c or '').strip().lower() for c in rows[0]]
    if 'source' not in head or 'result' not in head:
        return ['Attestation Log needs "Source" and "Result" columns']
    si, ri = head.index('source'), head.index('result')
    data = [r for r in rows[1:] if r and len(r) > si and str(r[si] or '').strip()]
    errs = []
    if len(data) != expected and not partial:
        errs.append(f'Attestation Log has {len(data)} source rows, but sources/{region}.md lists {expected}')
    for r in data:
        name = str(r[si]).strip()
        res = str(r[ri] or '').strip() if len(r) > ri else ''
        if not res:
            errs.append(f'no Result for "{name}"')
        elif any(b in res.lower() for b in BANNED_RESULTS):
            errs.append(f'banned Result for "{name}": {res[:80]}')
    errs += check_sources_used(wb)
    return errs


def check_sources_used(wb):
    """Sources Used sheet (template 2026-10-17): present, with a row for every Source # on the Weekly Sweep sheet."""
    if 'Sources Used' not in wb.sheetnames:
        return ['workbook has no "Sources Used" sheet (one row per source that produced a candidate)']
    rows = [r for r in wb['Sources Used'].iter_rows(values_only=True)]
    head = [str(c or '').strip().lower() for c in (rows[0] if rows else [])]
    if 'source #' not in head or 'source' not in head:
        return ['Sources Used needs "Source #" and "Source" columns']
    ni = head.index('source #')
    listed = {str(r[ni]).strip().lower() for r in rows[1:] if r and len(r) > ni and str(r[ni] or '').strip()}
    errs = []
    if 'Weekly Sweep' in wb.sheetnames:
        ws = [r for r in wb['Weekly Sweep'].iter_rows(values_only=True)]
        wh = [str(c or '').strip().lower() for c in (ws[0] if ws else [])]
        cand = [r for r in ws[1:] if r and len(r) > 1 and str(r[1] or '').strip()]
        if cand and not listed:
            errs.append('Sources Used is empty, but the Weekly Sweep sheet has candidates')
        if 'source #' in wh:
            si = wh.index('source #')
            for r in cand:
                v = str(r[si] or '').strip().lower() if len(r) > si else ''
                if not v:
                    errs.append(f'no Source # for candidate "{r[1]}"')
                elif v not in listed:
                    errs.append(f'Source # {v} (candidate "{r[1]}") is not on the Sources Used sheet')
        elif cand:
            errs.append('Weekly Sweep needs a "Source #" column')
    return errs


def self_test():
    ok = True
    tmp = tempfile.mkdtemp()
    try:
        shutil.copytree(os.path.join(HERE, 'sources'), os.path.join(tmp, 'sources'))
        shutil.copy(os.path.join(HERE, 'CLAUDE.md'), tmp)
        base = check(tmp)
        if base:
            print('self-test: real files are not clean, fix them first'); return False
        cases = [
            ('a rule section added to a region file', lambda t: t + '\n## Age categories\nTelluride uses different ages.\n'),
            ('a wrong Source count', lambda t: re.sub(r'(\| Source count \| )\d+', r'\g<1>999', t)),
            ('a missing header field', lambda t: re.sub(r'^\| Sweep window \|.*\n', '', t, flags=re.M)),
            ('a repeated source number', lambda t: t.replace('\n2. **', '\n1. **', 1)),
        ]
        target = os.path.join(tmp, 'sources', 'telluride.md')
        orig = open(target, encoding='utf-8').read()
        for name, mut in cases:
            open(target, 'w', encoding='utf-8').write(mut(orig))
            caught = bool(check(tmp))
            print(f'  {"PASS" if caught else "FAIL"}: catches {name}')
            ok &= caught
        open(target, 'w', encoding='utf-8').write(orig)
        os.remove(os.path.join(tmp, 'sources', 'napa.md'))
        caught = bool(check(tmp)); print(f'  {"PASS" if caught else "FAIL"}: catches a region in the table with no sources file'); ok &= caught
        # attestation completion check: a complete workbook passes, a short one and a banned result fail
        import openpyxl
        n = int(re.search(r'^\| Source count \| (\d+) \|', orig, re.M).group(1))
        def wb(rows, name, used=True, src_no='1'):
            w = openpyxl.Workbook(); ws = w.active; ws.title = 'Weekly Sweep'
            ws.append(['Decision', 'Event Name', 'Source #']); ws.append(['', 'Test event', src_no])
            if used:
                su = w.create_sheet('Sources Used'); su.append(['Source #', 'Source', 'Candidates (count)']); su.append(['1', 'Source 1', 1])
            a = w.create_sheet('Attestation Log'); a.append(['#', 'Source', 'Result'])
            for r in rows: a.append(r)
            path = os.path.join(tmp, name); w.save(path); return path
        full = [[i + 1, f'Source {i + 1}', 'DONE. fetched live list'] for i in range(n)]
        good = not check_attestation(wb(full, 'telluride_sweep_ok.xlsx'), root=HERE)
        print(f'  {"PASS" if good else "FAIL"}: accepts a complete attestation log'); ok &= good
        for label, rows in [('a missing source row', full[:-1]),
                            ('a "not reached" result', full[:-1] + [[n, 'Source N', 'not reached this pass']]),
                            ('a blank result', full[:-1] + [[n, 'Source N', '']])]:
            caught = bool(check_attestation(wb(rows, 'telluride_sweep_bad.xlsx'), root=HERE))
            print(f'  {"PASS" if caught else "FAIL"}: attestation check catches {label}'); ok &= caught
        for label, path in [('a missing Sources Used sheet', wb(full, 'telluride_sweep_nosu.xlsx', used=False)),
                            ('a candidate whose Source # is not on Sources Used', wb(full, 'telluride_sweep_badsrc.xlsx', src_no='7'))]:
            caught = bool(check_attestation(path, root=HERE))
            print(f'  {"PASS" if caught else "FAIL"}: attestation check catches {label}'); ok &= caught
        good = not check_attestation(wb(full[:1], 'telluride_sweep_partial.xlsx'), root=HERE, partial=True)
        print(f'  {"PASS" if good else "FAIL"}: --partial accepts a one-source review'); ok &= good
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print('self-test:', 'all passed' if ok else 'FAILED')
    return ok


if __name__ == '__main__':
    if '--self-test' in sys.argv:
        sys.exit(0 if self_test() else 1)
    if '--attestation' in sys.argv:
        x = sys.argv[sys.argv.index('--attestation') + 1]
        reg = sys.argv[sys.argv.index('--region') + 1] if '--region' in sys.argv else None
        errs = check_attestation(x, reg, partial='--partial' in sys.argv)
        for e in errs:
            print('  ✗', e)
        print(f'attestation: {"complete" if not errs else f"{len(errs)} problem(s)"} ({os.path.basename(x)})')
        sys.exit(1 if errs else 0)
    errs = check(HERE)
    for e in errs:
        print('  ✗', e)
    print(f'check_docs: {"clean" if not errs else f"{len(errs)} problem(s)"} ({len(os.listdir(os.path.join(HERE, "sources")))} region files)')
    sys.exit(1 if errs else 0)
